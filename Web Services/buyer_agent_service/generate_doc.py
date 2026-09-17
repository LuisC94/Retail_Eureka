import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set Margins
for sec in doc.sections:
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Title
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
title_run = title_p.add_run('Retail Eureka — Buyer Agent Web Service')
title_run.font.name = 'Calibri'
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(16, 44, 87)

# Subtitle
sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(16)
sub_run = sub_p.add_run('Manual de Protocolos de Comunicação REST API & Arquitetura Federada Multi-Tenant')
sub_run.font.name = 'Calibri'
sub_run.font.size = Pt(13)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(80, 80, 80)

# Section 1
h1 = doc.add_heading('1. Visão Geral & Arquitetura', level=1)
h1.paragraph_format.space_before = Pt(12)
h1.paragraph_format.space_after = Pt(6)

p1 = doc.add_paragraph(
    'O Web Service do Buyer Agent (Agente de Compras) é um microserviço autónomo construído com FastAPI '
    'e PyTorch, responsável pela otimização de encomendas diárias em cadeias de abastecimento de produtos perecíveis. '
    'O motor de decisão baseia-se em Aprendizagem por Reforço Profunda (PPO — Proximal Policy Optimization) com restrições '
    'físicas e modelo biológico de degradação de frutos (FruitModel2).'
)

doc.add_paragraph(
    'O serviço implementa isolamento multi-tenant por utilizador e cultura: os modelos treinados (.pth), scalers e '
    'metadados são guardados de forma estanque na pasta models_storage/user_{user_id}/culture_{culture_id}/.'
)

# Section 2 - Endpoints Summary Table
h2 = doc.add_heading('2. Sumário dos Endpoints REST', level=1)
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(6)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
headers = ['Método', 'Endpoint', 'Descrição', 'Porta']
for idx, text in enumerate(headers):
    hdr_cells[idx].text = text
    set_cell_background(hdr_cells[idx], '102C57')
    p = hdr_cells[idx].paragraphs[0]
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.bold = True

endpoints_data = [
    ('GET', '/health', 'Verificação de saúde do serviço', '8002'),
    ('GET', '/api/buyer/status', 'Consulta de existência e estado do modelo por user_id e culture_id', '8002'),
    ('POST', '/api/buyer/train', 'Treino da política PPO com histórico de vendas e previsões', '8002'),
    ('POST', '/api/buyer/decide', 'Inferência em tempo real para cálculo da encomenda ótima e baselines', '8002')
]

for m, ep, desc, port in endpoints_data:
    row_cells = table.add_row().cells
    row_cells[0].text = m
    row_cells[1].text = ep
    row_cells[2].text = desc
    row_cells[3].text = port
    set_cell_background(row_cells[0], 'F5F5F5')
    set_cell_background(row_cells[1], 'FFFFFF')
    set_cell_background(row_cells[2], 'FFFFFF')
    set_cell_background(row_cells[3], 'F5F5F5')

# Section 3 - Detailed Protocol
h3 = doc.add_heading('3. Detalhe dos Contratos de Comunicação', level=1)
h3.paragraph_format.space_before = Pt(14)
h3.paragraph_format.space_after = Pt(6)

doc.add_heading('3.1. POST /api/buyer/train (Treino de Política PPO)', level=2)
doc.add_paragraph('Parâmetros no corpo da mensagem (JSON):')
doc.add_paragraph(
    '• user_id (int/str): Identificador do utilizador autenticado.\n'
    '• culture_id (int/str): Identificador da cultura/subfamília de produto.\n'
    '• fruit_key (str, opcional): Preset biológico (ex: maca_gala, maca_fuji, kiwi_hayward, maca_golden, maca_reineta).\n'
    '• max_capacity (float, opcional): Limite físico de stock no armazém (default: 500.0).\n'
    '• episodes (int, opcional): Número de episódios de treino (default: 300).\n'
    '• train_data (array de objetos): Série histórica com real_value e prediction em cada dia.'
)

doc.add_heading('3.2. POST /api/buyer/decide (Tomada de Decisão Diária)', level=2)
doc.add_paragraph(
    'Permite que a plataforma Django ou qualquer cliente externo envie o estado de stock e o contexto de procura '
    'para obter a recomendação exata de compra em Kg e a comparação com baselines de mercado (DOS-3D, CNN Naive e Min-Max).'
)

doc.add_paragraph(
    'Vetor de Estado (17 variáveis processadas internamente):\n'
    '1-4. Stock físico classificado por RSL em 4 prateleiras (G0, G1, G2, G3)\n'
    '5. Encomendas em trânsito com chegada prevista para amanhã\n'
    '6-7. Previsão de procura para hoje e amanhã\n'
    '8-9. Vendas reais passadas (t-1 e t-2)\n'
    '10. Preço relativo de mercado (Z-score sobre janela de 15 dias)\n'
    '11-14. Componentes cíclicos temporais (seno e cosseno do dia da semana e do mês)\n'
    '15. Cobertura de stock normalizada\n'
    '16. Índice de urgência de stock a expirar\n'
    '17. Erro relativo da última previsão de procura'
)

# Section 4 - Heuristics comparison
h4 = doc.add_heading('4. Baselines Heurísticas para Comparação Transparente', level=1)
h4.paragraph_format.space_before = Pt(14)
h4.paragraph_format.space_after = Pt(6)

doc.add_paragraph(
    'A resposta do endpoint /api/buyer/decide inclui automaticamente a comparação com as 3 políticas tradicionais:\n'
    '• DOS 3-Day (Days of Supply): Garante cobertura para os próximos 3 dias de procura prevista.\n'
    '• CNN Naive Heuristic: Encomenda a diferença entre a procura de amanhã e o stock projetado de fim de dia.\n'
    '• Min-Max (s, S): Política clássica de ponto de encomenda.'
)

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "Protocolos_Comunicacao_BuyerAgent_WebService.docx")
doc.save(out_path)
print(f"Documento criado com sucesso em: {out_path}")
